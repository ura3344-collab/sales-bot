from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

import reports

ORDER = {"id": 7, "created": "2026-10-18 15:20", "status": "done", "name": "Ольга", "phone": "0501234567",
         "delivery": "Укрпошта, Львів", "payment": "💵 Накладений платіж", "comment": "", "total": 1040,
         "username": "olha", "rating": 5, "items": [{"title": "Фартух Latte (чорний)", "price": 520, "qty": 2}]}


def test_daily_report():
    since, until = datetime(2026, 10, 17, 16), datetime(2026, 10, 18, 16)
    cancelled = {**ORDER, "id": 8, "status": "cancelled", "total": 500}
    text = reports.daily_report([ORDER, cancelled], since, until)
    assert "17.10 16:00 — 18.10 16:00" in text
    assert "Замовлень: 2" in text and "Сума (без скасованих): 1040 грн" in text
    assert "№7 · 15:20 · Ольга · 1040 грн" in text
    assert "не було" in reports.daily_report([], since, until)


def test_orders_xlsx():
    ws = load_workbook(BytesIO(reports.orders_xlsx([ORDER]))).active
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["№"] == 7 and row["Сума, грн"] == 1040 and row["Оцінка"] == 5
    assert row["Товари"] == "Фартух Latte (чорний) × 2" and row["Telegram"] == "@olha"
